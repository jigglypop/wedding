import sys, os, json, datetime, hashlib
from pathlib import Path
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'weding-extract-deps'))
import openpyxl, xlrd
from PIL import Image, ImageOps, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
def serial(value):
    if isinstance(value, (datetime.datetime, datetime.date)): return value.isoformat()
    return value

books=[]
for p in sorted(ROOT.iterdir()):
    if p.suffix.lower() not in ['.xlsx', '.xls']: continue
    out={'id':'sheet-'+hashlib.sha256(p.name.encode()).hexdigest()[:12], 'filename':p.name, 'title':p.stem, 'sheets':[]}
    if p.suffix.lower()=='.xlsx':
        wb=openpyxl.load_workbook(p, data_only=False, read_only=True)
        cache=openpyxl.load_workbook(p, data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows=[]
            for ri, row in enumerate(ws.iter_rows(),start=1):
                cells=[]
                for c in row:
                    if c.value is None: continue
                    item={'cell':c.coordinate,'value':serial(c.value)}
                    if c.data_type=='f': item['cachedValue']=serial(cache[ws.title][c.coordinate].value)
                    cells.append(item)
                if cells: rows.append({'row':ri,'cells':cells})
            out['sheets'].append({'title':ws.title,'rows':rows,'maxRow':ws.max_row,'maxColumn':ws.max_column})
    else:
        wb=xlrd.open_workbook(p)
        for ws in wb.sheets():
            rows=[]
            for ri in range(ws.nrows):
                cells=[]
                for ci in range(ws.ncols):
                    cell=ws.cell(ri,ci)
                    if cell.value=='': continue
                    value=cell.value
                    if cell.ctype==xlrd.XL_CELL_DATE: value=xlrd.xldate.xldate_as_datetime(value,wb.datemode).isoformat()
                    cells.append({'cell':xlrd.formula.colname(ci)+str(ri+1),'value':serial(value)})
                if cells: rows.append({'row':ri+1,'cells':cells})
            out['sheets'].append({'title':ws.name,'rows':rows,'maxRow':ws.nrows,'maxColumn':ws.ncols})
        out['extractionNote']='XLS cached values extracted; original formulas remain in original file.'
    books.append(out)

(ROOT/'data').mkdir(exist_ok=True)
(ROOT/'data'/'local-spreadsheets.json').write_text(json.dumps(books,ensure_ascii=False,indent=2),encoding='utf-8')
images=[p for p in sorted(ROOT.iterdir()) if p.suffix.lower() in ['.jpg','.png','.webp']]
W,H=1600,360*4
canvas=Image.new('RGB',(W,H),'#f5f3ee')
draw=ImageDraw.Draw(canvas)
catalog=[]
for i,p in enumerate(images):
    im=Image.open(p); catalog.append({'filename':p.name,'width':im.width,'height':im.height})
    thumb=ImageOps.contain(im.convert('RGB'),(380,305))
    x=(i%4)*400+(400-thumb.width)//2; y=(i//4)*360+10
    canvas.paste(thumb,(x,y))
    draw.text(((i%4)*400+12,(i//4)*360+320),str(i+1)+' '+p.name,fill='#222222')
sheet=Path(os.environ['TEMP'])/'weding-material-contact-sheet.jpg';canvas.save(sheet,quality=88)
print(json.dumps({'books':[{'filename':b['filename'],'sheets':[{'title':s['title'],'rows':len(s['rows'])} for s in b['sheets']]} for b in books],'imageCatalog':catalog,'contactSheet':str(sheet)},ensure_ascii=False))
